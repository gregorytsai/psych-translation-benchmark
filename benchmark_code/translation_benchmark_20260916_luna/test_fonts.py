from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

fonts = [
    "Heiti TC",
    "Arial Unicode MS",
    "Hiragino Sans GB",
    "Songti TC",
    "STHeiti",
    "Apple LiGothic",
    "PingFang TC",
    "SimSun",
]
doc = Document()
for font in fonts:
    p = doc.add_paragraph()
    r = p.add_run(f"{font}  心理學與精神醫學  恐慌症  正體中文測試")
    r.font.name = font
    r.font.size = Pt(18)
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), font)
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), font)
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), font)
doc.save("/benchmark/translation_benchmark_20260916_luna/font_test.docx")
