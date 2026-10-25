from __future__ import annotations

import csv
import json
import os
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "心理學精神醫學翻譯模型評測報告.docx"
FONT = "Arial Unicode MS"
NAVY = "17365D"
BLUE = "DCE6F1"
PALE = "F3F6FA"
MID = "5B6573"
LIGHT_BORDER = "D9D9D9"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def remove_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_table_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_cell_borders(cell, color=LIGHT_BORDER, size="6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = tc_borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_width(cell, inches: float) -> None:
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, size=None, bold=None, color=None, italic=None) -> None:
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, size=9, color=MID)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_fonts.set(qn("w:eastAsia"), FONT)
    r_pr.extend([r_fonts, color, underline])
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def style_table(table, widths, header=True, font_size=9.2) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for r_idx, row in enumerate(table.rows):
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        if r_idx == 0 and header:
            set_repeat_table_header(row)
        for c_idx, cell in enumerate(row.cells):
            if c_idx < len(widths):
                set_cell_width(cell, widths[c_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_borders(cell)
            if r_idx == 0 and header:
                set_cell_shading(cell, NAVY)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, PALE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(1)
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.line_spacing = 1.08
                for run in paragraph.runs:
                    set_run_font(run, size=font_size, bold=(r_idx == 0 and header), color=("FFFFFF" if r_idx == 0 and header else "000000"))


def add_table(doc, headers, rows, widths, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = str(h)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    style_table(table, widths, font_size=font_size)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def add_body(doc, text: str, bold_lead: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.28
    if bold_lead:
        r = p.add_run(bold_lead)
        set_run_font(r, size=11, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=11)
    return p


def add_bullet(doc, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.17)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=10.7)
    return p


def clean_output(text: str) -> str:
    text = text.replace(" [end of text]", "").strip()
    return text


def main() -> None:
    scores = json.loads((ROOT / "scores.json").read_text(encoding="utf-8"))
    cases = [json.loads(line) for line in (ROOT / "inputs.jsonl").read_text(encoding="utf-8").splitlines() if line.strip()]
    image_case = json.loads((ROOT / "image_case.json").read_text(encoding="utf-8"))
    summary_rows = {}
    with (ROOT / "results_summary.csv").open(encoding="utf-8", newline="") as f:
        for row in csv.DictReader(f):
            summary_rows[row["model_id"]] = row

    models = [
        ("translategemma-27b", "TranslateGemma 27B", "Q8_0", ROOT.parent / "translategemma-27b-it-q8_0.gguf"),
        ("gemma-4-e4b", "Gemma 4 E4B", "Q4_K_XL", ROOT.parent / "gemma-4-E4B-it-qat-UD-Q4_K_XL.gguf"),
        ("gemma-4-12b", "Gemma 4 12B", "Q4_K_XL", ROOT.parent / "Gemma4 12b" / "gemma-4-12B-it-qat-UD-Q4_K_XL.gguf"),
        ("translategemma-12b", "TranslateGemma 12B", "Q8_0", ROOT.parent / "translategemma-12b-it.Q8_0.gguf"),
        ("translategemma-4b", "TranslateGemma 4B", "Q8_0", ROOT.parent / "translategemma-4b-it.Q8_0.gguf"),
    ]
    names = {m[0]: m[1] for m in models}

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    for style_name, size, before, after in (("Title", 24, 0, 12), ("Heading 1", 17, 14, 7), ("Heading 2", 13, 10, 5), ("Heading 3", 11.5, 8, 4)):
        st = styles[style_name]
        st.font.name = FONT
        st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        if style_name == "Title":
            style_p_pr = st._element.get_or_add_pPr()
            style_p_bdr = style_p_pr.find(qn("w:pBdr"))
            if style_p_bdr is not None:
                style_p_pr.remove(style_p_bdr)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("GGUF 本機翻譯評測  2026 年 9 月 16 日")
    set_run_font(run, size=8.5, color=MID)
    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph(style="Title")
    remove_paragraph_borders(title)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(88)
    title.add_run("心理學與精神醫學英文翻譯模型\n品質與效率評測報告")
    for run in title.runs:
        set_run_font(run, size=22, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(16)
    r = subtitle.add_run("TranslateGemma 27B 12B 4B 與 Gemma 4 E4B 12B")
    set_run_font(r, size=13, color=MID)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(34)
    r = meta.add_run("評測日期  2026 年 9 月 16 日\n評測平台  Apple Mac Studio M3 Ultra 256 GB\n評閱模型  Luna max")
    set_run_font(r, size=10.5, color=MID)
    lead = doc.add_paragraph()
    lead.paragraph_format.left_indent = Inches(0.55)
    lead.paragraph_format.right_indent = Inches(0.55)
    lead.paragraph_format.line_spacing = 1.35
    lead.paragraph_format.space_after = Pt(12)
    r = lead.add_run("結論先行  ")
    set_run_font(r, size=12, bold=True)
    r = lead.add_run("Gemma 4 E4B 提供最佳整體平衡。它的文字品質為 95.3 分，生成速度為 90.1 tokens/s，圖片翻譯為 98 分，主模型僅約 3.93 GiB。若只追求最高文字品質，TranslateGemma 27B 得分 99.3，但速度降至 21.8 tokens/s，檔案約 26.74 GiB。")
    set_run_font(r, size=12)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = note.add_run("本報告是小樣本本機比較，不構成臨床或醫療用途驗證。")
    set_run_font(r, size=10, italic=True, color=MID)
    doc.add_page_break()

    doc.add_heading("結論摘要", level=1)
    add_body(doc, "五個模型都能理解心理學與精神醫學內容，但在正體中文遵循、臨床術語與輸出潔淨度上差異明顯。品質與速度不能只按參數量推斷；本次 E4B 的整體表現優於兩個 12B 模型。")
    bullets = [
        "品質優先可選 TranslateGemma 27B。三段文字平均 99.3 分，臨床語義最穩定，但平均生成速度只有 21.8 tokens/s。",
        "日常使用首選 Gemma 4 E4B。文字品質 95.3 分、圖片翻譯 98 分、生成速度 90.1 tokens/s，兼顧容量與速度。",
        "Gemma 4 12B 的正文品質與 E4B 同為 95.3 分，但生成速度較低，且三段中都可看到控制標記或 thought channel 洩漏，需要後處理或更精確的模板設定。",
        "TranslateGemma 12B 的語義多數正確，但三段文字都輸出簡體中文，因此不適合直接交付台灣正體中文內容。",
        "TranslateGemma 4B 是最快的模型，達 100.7 tokens/s；然而它把 panic disorder 與 panic attack 誤譯為焦慮症與焦慮發作，不宜在無人工覆核下處理臨床文本。",
    ]
    for b in bullets:
        add_bullet(doc, b)

    doc.add_heading("整體比較", level=2)
    rec = {
        "translategemma-27b": "最高文字品質",
        "gemma-4-e4b": "最佳整體選擇",
        "gemma-4-12b": "高品質但需清理標記",
        "translategemma-12b": "需簡轉繁與覆核",
        "translategemma-4b": "僅適合低風險草稿",
    }
    overall_rows = []
    for mid, label, quant, path in models:
        s = summary_rows[mid]
        size = path.stat().st_size / (1024 ** 3)
        overall_rows.append([
            label,
            f"{size:.2f}",
            f"{float(s['text_avg_score_100']):.1f}",
            f"{float(s['text_avg_generation_tok_s']):.1f}",
            f"{float(s['text_avg_cold_wall_ms']) / 1000:.2f}",
            f"{float(s['image_ocr_score_100']):.1f}",
            f"{float(s['image_translation_score_100']):.0f}",
            rec[mid],
        ])
    add_table(doc, ["模型", "GiB", "文字品質", "生成速度", "文字總耗時秒", "圖片 OCR", "圖片翻譯", "建議"], overall_rows, [1.10, 0.42, 0.58, 0.62, 0.68, 0.55, 0.62, 1.68], font_size=8.0)

    doc.add_heading("評測範圍與方法", level=1)
    add_body(doc, "測試素材為三段自擬英文與一張本機合成的高對比臨床文字圖片。內容涵蓋恐慌症、行為活化、風險評估、SSRI 停藥症狀與血清素症候群。使用自擬素材可固定標準答案，也避免測試集洩漏與著作權問題。")
    add_body(doc, "文字測試使用相同提示、temperature 0、top k 1、seed 123、4096 context 與 256 token 上限。每次執行均重新載入模型，MTP 草稿模型一律停用。速度表分別呈現冷啟動總耗時、模型載入時間、提示處理速度與生成速度。")
    add_body(doc, "圖片測試要求模型先逐字轉錄 OCR，再翻譯為正體中文。圖片固定使用 1024 image tokens。OCR 採嚴格字元與內容檢查，圖片翻譯則比較臨床意義、藥名劑量、數字及正體中文表達。")
    add_body(doc, "Luna max 依一致量尺進行單一評閱。文字總分為 100 分，其中術語與領域準確度 30 分、語義忠實度 30 分、正體中文與流暢度 20 分、完整性與格式 20 分。模型標籤在評閱時可見，因此這不是正式盲評或多評者研究。")

    doc.add_heading("測試環境與模型", level=1)
    env_rows = [
        ["硬體", "Apple Mac Studio Mac15 14  Apple M3 Ultra  256 GB RAM  32 CPU cores"],
        ["作業系統", "Darwin 25.6.0 arm64"],
        ["文字執行器", "llama.cpp build 10679  Metal"],
        ["圖片執行器", "llama mtmd build 10679  Metal  強制 Gemma 相容模板"],
        ["抽樣", "temperature 0  top k 1  seed 123"],
        ["MTP", "全部停用以維持公平"],
    ]
    add_table(doc, ["項目", "設定"], env_rows, [1.35, 5.55], font_size=9.5)
    model_rows = []
    for mid, label, quant, path in models:
        model_rows.append([label, quant, f"{path.stat().st_size / (1024 ** 3):.2f} GiB", str(path)])
    add_table(doc, ["模型", "量化", "大小", "主模型路徑"], model_rows, [1.28, 0.78, 0.70, 4.14], font_size=8.4)

    doc.add_heading("文字翻譯結果", level=1)
    case_titles = {
        "panic_disorder": "恐慌症",
        "depression_behavioral_activation": "重度憂鬱症與行為活化",
        "ssri_pharmacology": "SSRI 與血清素症候群",
    }
    text_rows = []
    for mid, label, _, _ in models:
        case_scores = scores["text"][mid]["cases"]
        text_rows.append([
            label,
            str(case_scores["panic_disorder"]["total"]),
            str(case_scores["depression_behavioral_activation"]["total"]),
            str(case_scores["ssri_pharmacology"]["total"]),
            f"{scores['text'][mid]['average']:.1f}",
        ])
    add_table(doc, ["模型", "恐慌症", "行為活化", "SSRI", "平均"], text_rows, [2.15, 1.10, 1.20, 1.10, 1.10], font_size=9.5)
    add_body(doc, "TranslateGemma 27B 在三段中沒有重大術語錯誤。Gemma 4 E4B 與 12B 的正文也接近專業可用水準；E4B 的主要優勢是輸出乾淨，而 12B 的控制標記洩漏使交付品質下降。")
    add_body(doc, "TranslateGemma 12B 的簡體中文是本次最明顯的需求遵循失敗。TranslateGemma 4B 則有內容性錯誤，尤其把 panic 翻成 anxiety，會改變診斷與症狀名稱。")

    doc.add_heading("速度與效率", level=1)
    speed_rows = []
    for mid, label, _, _ in models:
        s = summary_rows[mid]
        speed_rows.append([
            label,
            f"{float(s['text_avg_load_ms']):.0f}",
            f"{float(s['text_avg_prompt_tok_s']):.1f}",
            f"{float(s['text_avg_generation_tok_s']):.1f}",
            f"{float(s['text_avg_cold_wall_ms']) / 1000:.2f}",
        ])
    add_table(doc, ["模型", "載入毫秒", "提示 tokens s", "生成 tokens s", "平均總耗時秒"], speed_rows, [2.25, 1.03, 1.22, 1.22, 1.20], font_size=9.2)
    add_body(doc, "TranslateGemma 4B 的生成速度最快，但品質不足以直接處理臨床文字。Gemma 4 E4B 的速度只低約 10.5 tokens/s，文字品質卻高出 19 分，因此具備最好的品質效率比。TranslateGemma 27B 的文字品質只比 E4B 高約 4 分，但生成速度不到四分之一，且檔案約大 6.8 倍。")

    doc.add_heading("圖片翻譯結果", level=1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(ROOT / "image_clinical_note.png"), width=Inches(6.45))
    cap = doc.add_paragraph("圖片測試樣本  本機合成的精神科追蹤文字")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        set_run_font(run, size=9, italic=True, color=MID)
    image_note_zh = {
        "TranslateGemma 27B": "OCR 完全正確；「複診」與「追蹤」意義接近，藥名與數字皆保留。",
        "TranslateGemma 12B": "FOUR WEEKS 後的逗號誤為句點；翻譯增加請求與命令語氣。",
        "TranslateGemma 4B": "OCR 完全正確；panic attack 誤為焦慮發作，標題與語序較生硬。",
        "Gemma 4 E4B": "OCR 完全正確；譯文自然，數字、劑量與藥名皆保留。",
        "Gemma 4 12B": "OCR 完全正確；「每週」重複一次，Sertraline 保留拉丁字母。",
    }
    image_rows = []
    for mid, label, _, _ in models:
        item = scores["image"][mid]
        image_rows.append([
            label,
            "是" if item["ocr_exact_match"] else "否",
            f"{item['ocr_score']:.1f}",
            str(item["translation_score"]),
            f"{item['wall_ms'] / 1000:.2f}",
            image_note_zh[label],
        ])
    image_table = add_table(doc, ["模型", "OCR 完全相符", "OCR", "翻譯", "耗時秒", "主要觀察"], image_rows, [1.35, 0.78, 0.55, 0.55, 0.62, 3.02], font_size=8.4)
    for row in image_table.rows:
        keep_table_row_together(row)
    add_body(doc, "四個模型的 OCR 與原文逐字相符。TranslateGemma 12B 只把 FOUR WEEKS 後的逗號辨識為句點，核心內容、藥名、劑量與數字均正確。圖片翻譯以 Gemma 4 E4B 得分最高；TranslateGemma 4B 再次把 panic attack 誤譯為焦慮發作。")

    doc.add_heading("圖片完整成果", level=2)
    for mid, label, _, _ in models:
        output_path = ROOT / "raw_image_metal" / f"{mid}__image_clinical_note.output"
        output = output_path.read_text(encoding="utf-8", errors="replace").strip()
        doc.add_heading(label, level=3)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.right_indent = Inches(0.12)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(output)
        set_run_font(r, size=9.3)

    doc.add_heading("逐段完整翻譯成果", level=1)
    add_body(doc, "以下表格保留各模型的實際文字輸出，只移除 llama.cpp 測試工具加入的 end of text 結尾標記。模型自行洩漏的 channel 或 thought 標記保留，以反映真實交付品質。")
    for case in cases:
        cid = case["id"]
        doc.add_heading(case_titles[cid], level=2)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run("英文原文  ")
        set_run_font(r, size=9.5, bold=True)
        r = p.add_run(case["english"])
        set_run_font(r, size=9.5)
        rows = []
        for mid, label, _, _ in models:
            out_path = ROOT / "raw_text_metal" / f"{mid}__{cid}.stdout"
            translation = clean_output(out_path.read_text(encoding="utf-8", errors="replace"))
            score = scores["text"][mid]["cases"][cid]
            rows.append([label, translation, f"{score['total']}\n{score['notes']}"])
        add_table(doc, ["模型", "完整翻譯輸出", "分數與觀察"], rows, [1.20, 4.25, 1.45], font_size=8.6)

    doc.add_heading("限制與建議", level=1)
    limitations = [
        "樣本只有三段文字與一張圖片，足以辨識明顯錯誤，但不能代表所有精神醫學次專科、文件類型或長篇內容。",
        "品質分數由單一 Luna max 評閱者產生，且模型標籤可見。後續正式採購或部署應加入至少兩名具臨床翻譯經驗的人類評閱者，並採盲評與一致性分析。",
        "不同模型使用不同量化。TranslateGemma 是 Q8 0，Gemma 4 是 Q4 K XL；結果代表資料夾內現有檔案，不是原始模型架構的純比較。",
        "圖片 CLI 使用相容的 Gemma 模板，而不是 TranslateGemma 官方需要結構化語言代碼的原生 Jinja 輸入。圖片結果代表目前 llama.cpp GGUF 工作流程。",
        "Gemma 4 12B 的控制標記可能透過新版 chat template 或輸出清理解決。TranslateGemma 12B 的簡體輸出也可後轉為正體，但簡繁轉換不能修正術語或語義錯誤。",
        "所有醫療相關翻譯仍需人工覆核，特別是診斷、風險描述、藥名、劑量、否定詞與不確定性。",
    ]
    for item in limitations:
        add_bullet(doc, item)

    doc.add_heading("部署建議", level=2)
    deploy_rows = [
        ["大量草稿與人工覆核", "Gemma 4 E4B", "速度快、容量小、文字與圖片品質穩定"],
        ["高品質文字翻譯", "TranslateGemma 27B", "品質最高但速度與容量成本較高"],
        ["圖片為主要輸入", "Gemma 4 E4B", "本次圖片翻譯最高分且 OCR 完全正確"],
        ["低資源設備", "TranslateGemma 4B", "只建議低風險草稿，必須人工校對臨床術語"],
        ["正體中文直接交付", "不建議 TranslateGemma 12B", "目前穩定輸出簡體中文"],
    ]
    add_table(doc, ["情境", "建議模型", "理由"], deploy_rows, [1.55, 1.55, 3.80], font_size=9.2)

    doc.add_heading("原始資料與重現方法", level=1)
    paths = [
        ("評測根目錄", str(ROOT)),
        ("英文測試素材", str(ROOT / "inputs.jsonl")),
        ("圖片測試定義", str(ROOT / "image_case.json")),
        ("文字原始輸出", str(ROOT / "raw_text_metal")),
        ("圖片原始輸出", str(ROOT / "raw_image_metal")),
        ("評分", str(ROOT / "scores.json")),
        ("摘要資料", str(ROOT / "results_summary.csv")),
        ("重現腳本", str(ROOT / "run_text_benchmark.sh")),
    ]
    add_table(doc, ["項目", "路徑"], paths, [1.55, 5.35], font_size=8.6)
    add_body(doc, "官方資料指出 TranslateGemma 是以 Gemma 3 為基礎的 55 語言翻譯模型，支援文字與圖片輸入；Gemma 4 則是多模態模型系列。本報告的數值來自本機 GGUF 實測，不採用官方宣稱作為品質分數。")
    p = doc.add_paragraph()
    r = p.add_run("TranslateGemma 官方模型卡  ")
    set_run_font(r, size=9.5, bold=True)
    add_hyperlink(p, "Hugging Face", "https://huggingface.co/google/translategemma-27b-it")
    p.add_run("    ")
    r = p.add_run("Gemma 4 官方模型卡  ")
    set_run_font(r, size=9.5, bold=True)
    add_hyperlink(p, "Google AI for Developers", "https://ai.google.dev/gemma/docs/core/model_card_4")

    doc.core_properties.title = "心理學與精神醫學英文翻譯模型品質與效率評測報告"
    doc.core_properties.subject = "GGUF 翻譯模型本機評測"
    doc.core_properties.author = "Codex 與 Luna max"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
